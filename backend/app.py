from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from werkzeug.utils import secure_filename
from pathlib import Path

from analyzer import analyze_resume, match_job_description

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

BASE = Path(__file__).resolve().parent.parent
FRONTEND = BASE / "frontend"
UPLOADS = BASE / "uploads"
UPLOADS.mkdir(exist_ok=True)

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ALLOWED = {"pdf", "docx", "txt"}

def extract_text(path):
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        if PdfReader is None:
            raise RuntimeError("PDF support is unavailable.")
        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        if Document is None:
            raise RuntimeError("DOCX support is unavailable.")
        doc = Document(str(path))
        text = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                text.append(" ".join(c.text for c in row.cells))
        return "\n".join(text)
    raise ValueError("Unsupported file type.")

@app.get("/api/health")
def health():
    return jsonify({"status": "online", "app": "ResumeX AI V2"})

@app.post("/api/analyze")
def analyze():
    if "resume" not in request.files:
        return jsonify({"success": False, "message": "Please select a resume."}), 400
    f = request.files["resume"]
    if not f.filename:
        return jsonify({"success": False, "message": "Please select a resume."}), 400
    ext = Path(f.filename).suffix.lower().lstrip(".")
    if ext not in ALLOWED:
        return jsonify({"success": False, "message": "Use PDF, DOCX or TXT."}), 400

    path = UPLOADS / secure_filename(f.filename)
    f.save(path)
    try:
        text = extract_text(path).strip()
        if not text:
            return jsonify({"success": False, "message": "No readable text found in this resume."}), 400
        result = analyze_resume(text)
        return jsonify({"success": True, "analysis": result})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@app.post("/api/match")
def match():
    data = request.get_json(silent=True) or {}
    resume_text = data.get("resume_text", "").strip()
    job_text = data.get("job_description", "").strip()
    if not resume_text or not job_text:
        return jsonify({"success": False, "message": "Resume and job description are required."}), 400
    return jsonify({"success": True, "match": match_job_description(resume_text, job_text)})

@app.route("/")
def home():
    return send_from_directory(FRONTEND, "index.html")

@app.route("/<path:path>")
def static_files(path):
    return send_from_directory(FRONTEND, path)

if __name__ == "__main__":
    print("=" * 55)
    print("             RESUMEX AI V2")
    print("       PROFESSIONAL RESUME ANALYZER")
    print("=" * 55)
    print("http://127.0.0.1:5000")
    print("=" * 55)
    app.run(host="127.0.0.1", port=5000, debug=False)
